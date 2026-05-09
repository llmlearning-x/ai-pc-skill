"""
Local Privacy Inspector - File Text Extractor
支持从多种文件类型中提取纯文本内容
"""

import os
import csv
import json
from io import StringIO
from typing import Optional, Tuple

# 支持的文件类型
SUPPORTED_EXTENSIONS = {
    '.txt', '.md', '.env', '.json', '.yaml', '.yml', '.csv',
    '.docx', '.pdf', '.xlsx'
}


def extract_text(file_path: str) -> Tuple[Optional[str], Optional[str]]:
    """
    从文件中提取文本内容
    
    Args:
        file_path: 文件路径
    
    Returns:
        (text_content, error_message)
        成功时 error_message 为 None
        失败时 text_content 为 None
    """
    if not os.path.exists(file_path):
        return None, f"文件不存在: {file_path}"
    
    if not os.path.isfile(file_path):
        return None, f"路径不是文件: {file_path}"
    
    file_name = os.path.basename(file_path)
    ext = os.path.splitext(file_path)[1].lower()
    
    # 特殊处理以点开头的文件（如 .env）
    if file_name.startswith('.') and '.' in file_name[1:]:
        ext = os.path.splitext(file_name)[1].lower()
    elif file_name.startswith('.'):
        ext = file_name.lower()
    
    if ext not in SUPPORTED_EXTENSIONS:
        return None, f"不支持的文件类型: {ext}"
    
    file_size = os.path.getsize(file_path)
    if file_size > 10 * 1024 * 1024:  # 10MB
        return None, f"文件超过10MB限制: {file_size} 字节"
    
    try:
        if ext in {'.txt', '.md', '.env'}:
            return _extract_plain_text(file_path), None
        elif ext == '.json':
            return _extract_json(file_path), None
        elif ext in {'.yaml', '.yml'}:
            return _extract_plain_text(file_path), None
        elif ext == '.csv':
            return _extract_csv(file_path), None
        elif ext == '.docx':
            return _extract_docx(file_path), None
        elif ext == '.pdf':
            return _extract_pdf(file_path), None
        elif ext == '.xlsx':
            return _extract_xlsx(file_path), None
        else:
            return None, f"暂不支持的文件格式: {ext}"
    except Exception as e:
        return None, f"文件解析失败: {str(e)}"


def _extract_plain_text(file_path: str) -> str:
    """直接读取文本文件"""
    encodings = ['utf-8', 'utf-8-sig', 'gbk', 'gb2312', 'latin-1']
    for encoding in encodings:
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码")


def _extract_json(file_path: str) -> str:
    """提取 JSON 文件的文本内容（扁平化）"""
    text = _extract_plain_text(file_path)
    try:
        data = json.loads(text)
        return _flatten_json(data)
    except json.JSONDecodeError:
        # 如果解析失败，返回原始文本
        return text


def _flatten_json(data, prefix='') -> str:
    """将 JSON 对象扁平化为文本"""
    lines = []
    if isinstance(data, dict):
        for key, value in data.items():
            new_prefix = f"{prefix}.{key}" if prefix else key
            if isinstance(value, (dict, list)):
                lines.append(_flatten_json(value, new_prefix))
            else:
                lines.append(f"{new_prefix}: {value}")
    elif isinstance(data, list):
        for i, item in enumerate(data):
            new_prefix = f"{prefix}[{i}]"
            if isinstance(item, (dict, list)):
                lines.append(_flatten_json(item, new_prefix))
            else:
                lines.append(f"{new_prefix}: {item}")
    else:
        lines.append(f"{prefix}: {data}")
    return '\n'.join(lines)


def _extract_csv(file_path: str) -> str:
    """提取 CSV 文件的文本内容"""
    text = _extract_plain_text(file_path)
    lines = []
    try:
        reader = csv.reader(StringIO(text))
        for row in reader:
            lines.append(', '.join(row))
    except Exception:
        # 如果 CSV 解析失败，返回原始文本
        return text
    return '\n'.join(lines)


def _extract_docx(file_path: str) -> str:
    """提取 Word 文档 (.docx) 的段落文本"""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("解析 .docx 需要 python-docx，请执行: pip install python-docx")
    
    doc = Document(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)
    return '\n'.join(paragraphs)


def _extract_pdf(file_path: str) -> str:
    """提取 PDF 文件的文本内容"""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError("解析 .pdf 需要 pypdf，请执行: pip install pypdf")
    
    reader = PdfReader(file_path)
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages.append(f"--- 第 {i + 1} 页 ---\n{text.strip()}")
    return '\n\n'.join(pages)


def _extract_xlsx(file_path: str) -> str:
    """提取 Excel 表格 (.xlsx) 的单元格文本"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("解析 .xlsx 需要 openpyxl，请执行: pip install openpyxl")
    
    wb = load_workbook(file_path, data_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        lines.append(f"--- 工作表: {sheet_name} ---")
        for row in sheet.iter_rows(values_only=True):
            row_text = ', '.join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                lines.append(row_text)
    return '\n'.join(lines)


def get_file_type(file_path: str) -> str:
    """获取文件类型"""
    ext = os.path.splitext(file_path)[1].lower()
    return ext.lstrip('.') if ext else 'unknown'
